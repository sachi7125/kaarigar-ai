"""Per-marketplace exports (Day 7): built on demand from one listing, columns
from the profile files, anything we don't collect flagged rather than guessed."""
from __future__ import annotations

import io
import json

import pytest
from openpyxl import load_workbook

from app.db.models import Artisan, Listing, ListingPhoto
from app.services import marketplace_export as mx


def _listing(db, **overrides) -> Listing:
    artisan = Artisan(phone=f"+91{db.query(Artisan).count():010d}", verified=True)
    db.add(artisan)
    db.commit()
    fields = dict(artisan_id=artisan.id, category="clay pottery", material="clay", size_class="medium",
                  title_en="Blue clay matki", title_hi="नीली मटकी",
                  description_en="A hand-thrown pot.", description_hi="हाथ से बनी मटकी।",
                  price_inr=850, stock_type="batch", total_count=5, remaining_count=4)
    fields.update(overrides)
    listing = Listing(**fields)
    db.add(listing)
    db.commit()
    db.refresh(listing)
    return listing


def _sheet(body: bytes):
    wb = load_workbook(io.BytesIO(body))
    ws = wb.worksheets[0]
    headers = [c.value for c in ws[1]]
    return wb, ws, dict(zip(headers, [c.value for c in ws[2]]))


def test_the_four_marketplaces_are_offered():
    assert {m["id"] for m in mx.available_marketplaces()} == {"amazon", "flipkart", "gem", "ondc"}


@pytest.mark.parametrize("marketplace", ["amazon", "flipkart", "gem"])
def test_spreadsheet_follows_the_profile_and_flags_what_we_dont_collect(db_session, marketplace):
    listing = _listing(db_session)
    body, media_type, filename = mx.build_export(listing, marketplace)
    assert media_type == mx.XLSX_MEDIA_TYPE
    assert filename == f"kaarigar_{marketplace}_{listing.id}.xlsx"

    wb, ws, row = _sheet(body)
    profile = mx.load_profile(marketplace)
    assert list(row) == [c["header"] for c in profile["columns"]]
    for i, col in enumerate(profile["columns"], start=1):
        cell = ws.cell(row=2, column=i)
        if "not_collected" in col:
            assert str(cell.value).startswith("not collected — ")
            assert cell.fill.fgColor.rgb.endswith("FFF3CD")

    about = " ".join(str(c.value) for r in wb["About this file"].iter_rows() for c in r if c.value)
    assert profile["verified_on"] in about and profile["sources"][0] in about


def test_known_values_land_in_the_right_columns(db_session):
    listing = _listing(db_session)
    _, _, amazon = _sheet(mx.build_export(listing, "amazon")[0])
    assert amazon["item_sku"] == listing.id and amazon["item_name"] == "Blue clay matki"
    assert amazon["standard_price"] == 850 and amazon["quantity"] == 4
    assert amazon["material_type"] == "clay" and amazon["country_of_origin"] == "India"
    assert amazon["hsn_code"].startswith("not collected")

    _, _, flipkart = _sheet(mx.build_export(listing, "flipkart")[0])
    assert flipkart["Your Selling Price"] == 850 and flipkart["MRP"].startswith("not collected")

    _, _, gem = _sheet(mx.build_export(listing, "gem")[0])
    assert gem["Unit of Measurement"] == "Piece" and gem["Offer Quantity"] == 4


def test_a_set_is_one_unit_of_n_pieces(db_session):
    listing = _listing(db_session, stock_type="unique", total_count=1, remaining_count=1,
                       price_unit="set", pack_size=12)
    _, _, gem = _sheet(mx.build_export(listing, "gem")[0])
    assert gem["Unit of Measurement"] == "Set" and gem["Offer Quantity"] == 1
    _, _, flipkart = _sheet(mx.build_export(listing, "flipkart")[0])
    assert flipkart["Pack Of"] == 12


def test_photo_url_only_when_there_is_a_photo(db_session):
    without = _listing(db_session)
    _, _, row = _sheet(mx.build_export(without, "amazon")[0])
    assert row["main_image_url"] == "not collected — no photo on this listing"

    with_photo = _listing(db_session)
    db_session.add(ListingPhoto(listing_id=with_photo.id, data=b"jpeg"))
    db_session.commit()
    db_session.refresh(with_photo)
    _, _, row = _sheet(mx.build_export(with_photo, "amazon")[0])
    assert row["main_image_url"].endswith(f"/media/listing/{with_photo.id}")
    assert row["other_image_url1"].startswith("not collected")     # a plain photo has no close-up

    enhanced = _listing(db_session)
    db_session.add(ListingPhoto(listing_id=enhanced.id, kind="enhanced", data=b"jpeg", texture_data=b"jpeg"))
    db_session.commit()
    db_session.refresh(enhanced)
    _, _, row = _sheet(mx.build_export(enhanced, "amazon")[0])
    assert row["other_image_url1"].endswith(f"/media/listing/{enhanced.id}/texture")
    images = json.loads(mx.build_export(enhanced, "ondc")[0])["item"]["descriptor"]["images"]
    assert len(images) == 2 and images[1].endswith("/texture")


def test_ondc_item_json(db_session):
    listing = _listing(db_session)
    body, media_type, filename = mx.build_export(listing, "ondc")
    assert media_type == "application/json" and filename.endswith(".json")
    doc = json.loads(body)
    item = doc["item"]
    assert item["id"] == listing.id
    assert item["descriptor"]["name"] == "Blue clay matki"
    assert item["descriptor"]["images"] == []            # no photo: empty, not a broken URL
    assert item["price"] == {"currency": "INR", "value": "850.00", "maximum_value": "850.00"}
    assert item["quantity"]["available"]["count"] == 4    # a number, not the string "4"
    assert item["category_id"] == "Home & Decor"
    assert item["@ondc/org/returnable"] is None
    assert doc["fields_to_fill"] and doc["sources"]


def test_ondc_leaves_the_category_to_the_seller_app_outside_home_decor(db_session):
    listing = _listing(db_session, category="leather footwear", material="leather")
    assert json.loads(mx.build_export(listing, "ondc")[0])["item"]["category_id"] is None


@pytest.mark.parametrize("name", ["myntra", "../gem", "", "GEM"])
def test_unknown_marketplace(db_session, name):
    listing = _listing(db_session)
    with pytest.raises(mx.UnknownMarketplace):
        mx.build_export(listing, name)
