"""Per-marketplace export files, built on demand from one listing (Day 7).

Each marketplace is a small JSON profile in data/reference/marketplace_profiles/
mapping our listing fields onto its columns. Nothing is generated ahead of time
or stored: the endpoint builds the file in memory when she taps "Export" for a
marketplace and streams it back. Adding or correcting a marketplace is a data
edit, not a code change.

Two honesty rules, extending decision D22:
- Amazon, Flipkart and GeM don't publish one fixed upload format; templates are
  per-category and downloadable only inside a seller account. Each profile is
  built from publicly documented fields, dated, and says so inside the file.
- A field we don't collect is written as an explicit "not collected — ..." note
  and highlighted, never guessed: no invented brand, HSN code, barcode, MRP or
  return policy.
"""
from __future__ import annotations

import io
import json
import re
from functools import lru_cache

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.config import PUBLIC_URL, REPO_ROOT
from app.db.models import Listing
from pipelines.voice.describe import normalize_category

PROFILES_DIR = REPO_ROOT / "data" / "reference" / "marketplace_profiles"
NOT_COLLECTED = "not collected"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_PLACEHOLDER = re.compile(r"\{(\w+)\}")
_TO_FILL = PatternFill("solid", fgColor="FFF3CD")

# ONDC's retail item category_id values include "Home & Decor". Textiles,
# footwear, jewellery and bags sit in domains whose codes the seller app
# assigns, so those are left for it rather than guessed here.
_ONDC_HOME_DECOR = frozenset({
    "clay pottery", "terracotta diya", "brass idol", "dhokra figurine",
    "copper vessel", "bamboo basket", "wooden toy", "madhubani painting",
})


class UnknownMarketplace(LookupError):
    pass


@lru_cache(maxsize=None)
def load_profile(marketplace: str) -> dict:
    if not re.fullmatch(r"[a-z0-9_]+", marketplace or ""):
        raise UnknownMarketplace(marketplace)
    path = PROFILES_DIR / f"{marketplace}.json"
    if not path.exists():
        raise UnknownMarketplace(marketplace)
    return json.loads(path.read_text(encoding="utf-8"))


def available_marketplaces() -> list[dict]:
    out = []
    for path in sorted(PROFILES_DIR.glob("*.json")):
        profile = load_profile(path.stem)
        out.append({"id": path.stem, "name": profile["display_name"], "format": profile["format"]})
    return out


def listing_fields(listing: Listing) -> dict:
    category = (listing.category or "").strip()
    material = (listing.material or "").strip()
    per_set = listing.price_unit == "set"
    has_photo = listing.photo is not None or bool(listing.image_path)
    has_closeup = listing.photo is not None and listing.photo.kind == "enhanced"
    return {
        "listing_id": listing.id,
        "title_en": listing.title_en,
        "title_hi": listing.title_hi,
        "description_en": listing.description_en,
        "description_hi": listing.description_hi,
        "category": category,
        "material": material,
        "size_class": (listing.size_class or "").strip(),
        "price_inr": round(float(listing.price_inr), 2),
        "price_inr_str": f"{float(listing.price_inr):.2f}",
        "remaining_count": listing.remaining_count,
        "uom": "Set" if per_set else "Piece",
        "pack_of": listing.pack_size if per_set else 1,
        "keywords": ", ".join(x for x in ("handmade", category, material) if x),
        "features": "; ".join(x for x in (f"Material: {material}" if material else "",
                                            "Handmade in India") if x),
        "image_url": f"{PUBLIC_URL}/media/listing/{listing.id}" if has_photo else None,
        "texture_url": f"{PUBLIC_URL}/media/listing/{listing.id}/texture" if has_closeup else None,
        "listing_url": f"{PUBLIC_URL}/l/{listing.id}",
        "storefront_url": f"{PUBLIC_URL}/s/{listing.artisan_id}",
        "ondc_category_id": ("Home & Decor"
                             if normalize_category(category, "", [material]) in _ONDC_HOME_DECOR else None),
    }


def _fill(template, fields: dict):
    """Substitute {field} placeholders. A string that is exactly one placeholder
    keeps the field's own type (a count stays an int) and becomes None when the
    field is empty; None entries are dropped from lists."""
    if isinstance(template, str):
        whole = _PLACEHOLDER.fullmatch(template)
        if whole:
            value = fields.get(whole.group(1))
            return None if value in (None, "") else value
        return _PLACEHOLDER.sub(lambda m: str(fields.get(m.group(1)) or ""), template)
    if isinstance(template, list):
        return [v for v in (_fill(t, fields) for t in template) if v is not None]
    if isinstance(template, dict):
        return {k: _fill(v, fields) for k, v in template.items()}
    return template


def _column_value(col: dict, fields: dict) -> tuple[object, str | None, bool]:
    """(cell value, note for the About sheet, still to fill?)"""
    if "not_collected" in col:
        return f"{NOT_COLLECTED} — {col['not_collected']}", col["not_collected"], True
    template = col["value"]
    if any(fields.get(name) in (None, "") for name in _PLACEHOLDER.findall(template)):
        why = col.get("if_empty", "not captured for this listing")
        return f"{NOT_COLLECTED} — {why}", why, True
    return _fill(template, fields), col.get("note"), False


def build_xlsx(profile: dict, fields: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = profile.get("sheet_name", "Listing")
    cells = [(col["header"], *_column_value(col, fields)) for col in profile["columns"]]
    ws.append([header for header, _, _, _ in cells])
    ws.append([value for _, value, _, _ in cells])
    for i, (_, _, _, to_fill) in enumerate(cells, start=1):
        ws.cell(row=1, column=i).font = Font(bold=True)
        if to_fill:
            ws.cell(row=2, column=i).fill = _TO_FILL
    for col in ws.columns:
        width = max(len(str(c.value)) for c in col if c.value is not None)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 12), 60)

    about = wb.create_sheet("About this file")
    about.append(["KaarigarAI export for", profile["display_name"]])
    about.append(["Listing page", fields["listing_url"]])
    about.append(["Columns last checked", profile["verified_on"]])
    about.append([])
    for line in profile["about"]:
        about.append([line])
    notes = [(header, note, to_fill) for header, _, note, to_fill in cells if note]
    if notes:
        about.append([])
        about.append(["Column", "What to do"])
        for header, note, to_fill in notes:
            about.append([header, ("Fill in: " if to_fill else "Check: ") + note])
    about.append([])
    about.append(["Sources"])
    for source in profile["sources"]:
        about.append([source])
    about.column_dimensions["A"].width = 28
    about.column_dimensions["B"].width = 90

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_json(profile: dict, fields: dict) -> bytes:
    body = {
        "about": profile["about"],
        "listing_page": fields["listing_url"],
        "fields_last_checked": profile["verified_on"],
        "fields_to_fill": profile.get("fields_to_fill", []),
        "sources": profile["sources"],
        "item": _fill(profile["template"], fields),
    }
    return json.dumps(body, ensure_ascii=False, indent=2).encode("utf-8")


def build_export(listing: Listing, marketplace: str) -> tuple[bytes, str, str]:
    """(file bytes, media type, filename). Raises UnknownMarketplace."""
    profile = load_profile(marketplace)
    fields = listing_fields(listing)
    if profile["format"] == "xlsx":
        return build_xlsx(profile, fields), XLSX_MEDIA_TYPE, f"kaarigar_{marketplace}_{listing.id}.xlsx"
    if profile["format"] == "json":
        return build_json(profile, fields), "application/json", f"kaarigar_{marketplace}_{listing.id}.json"
    raise ValueError(f"unsupported export format {profile['format']!r} for {marketplace}")
